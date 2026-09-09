import io
import json
import pytest
import docx
from fastapi.testclient import TestClient

from src.db.connection import get_connection
from src.db.migrations import init_db
from src.db.models import DocumentModel, NodeModel, ProjectModel
from src.main import app
from src.parsers.docx_parser import parse_docx
from src.parsers.text_parser import parse_markdown, parse_plain_text
from src.parsers.pdf_parser import parse_pdf_fallback
from src.services.hierarchy_service import (
    detach_selection_to_header,
    merge_nodes,
    promote_node_to_header,
    split_node,
)
from src.services.project_manager import register_custom_connection

client = TestClient(app)

@pytest.fixture
def memory_project():
    project_id = "test_ingestion_proj"
    conn = get_connection(":memory:")
    init_db(conn)

    with conn:
        conn.execute(
            "INSERT INTO projects (id, name, llm_model, embedding_model) VALUES (?, 'Test Project', 'llm', 'emb');",
            (project_id,),
        )
        conn.execute(
            "INSERT INTO documents (id, project_id, filename, file_type, order_index) VALUES ('doc_1', ?, 'doc1.docx', 'docx', 0);",
            (project_id,),
        )

    register_custom_connection(project_id, conn)
    yield (project_id, conn)
    conn.close()

def test_docx_parser_heading_hierarchy():
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("First paragraph describing architecture.")
    doc.add_heading("Subsystem A", level=2)
    doc.add_paragraph("Details of subsystem A.")
    doc.add_heading("Subsystem B", level=2)
    doc.add_paragraph("Details of subsystem B.")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    nodes = parse_docx(file_stream, document_id="doc_test")
    assert len(nodes) == 6

    # Heading 1: Architecture Overview
    h1 = nodes[0]
    assert h1.node_type == "header"
    assert h1.text_content == "Architecture Overview"
    assert h1.parent_id is None

    # Paragraph 1 linked to Heading 1
    p1 = nodes[1]
    assert p1.node_type == "paragraph"
    assert p1.parent_id == h1.id

    # Heading 2: Subsystem A linked to Heading 1
    h2_a = nodes[2]
    assert h2_a.node_type == "header"
    assert h2_a.text_content == "Subsystem A"
    assert h2_a.parent_id == h1.id

    # Paragraph 2 linked to Subsystem A
    p2_a = nodes[3]
    assert p2_a.parent_id == h2_a.id

    # Heading 2: Subsystem B linked to Heading 1
    h2_b = nodes[4]
    assert h2_b.parent_id == h1.id

    # Paragraph 3 linked to Subsystem B
    p2_b = nodes[5]
    assert p2_b.parent_id == h2_b.id

def test_markdown_parser_hierarchy():
    md = """# Title Header
Intro paragraph line 1.
Intro paragraph line 2.

## Section 1
Section 1 text.

### Detail Subsection
Nested detail paragraph.
"""
    nodes = parse_markdown(md, document_id="doc_md")
    assert len(nodes) == 6

    assert nodes[0].node_type == "header"
    assert nodes[0].text_content == "Title Header"
    assert nodes[0].parent_id is None

    assert nodes[1].node_type == "paragraph"
    assert nodes[1].parent_id == nodes[0].id

    assert nodes[2].node_type == "header"
    assert nodes[2].text_content == "Section 1"
    assert nodes[2].parent_id == nodes[0].id

    assert nodes[3].node_type == "paragraph"
    assert nodes[3].parent_id == nodes[2].id

    assert nodes[4].node_type == "header"
    assert nodes[4].text_content == "Detail Subsection"
    assert nodes[4].parent_id == nodes[2].id

    assert nodes[5].node_type == "paragraph"
    assert nodes[5].parent_id == nodes[4].id

def test_plain_text_parser():
    txt = "Paragraph one text.\n\nParagraph two with more details.\n\nParagraph three."
    nodes = parse_plain_text(txt, document_id="doc_txt")
    assert len(nodes) == 3
    assert [n.order_index for n in nodes] == [0, 1, 2]
    assert nodes[0].text_content == "Paragraph one text."

def test_manual_split_maintains_order_indices_and_status(memory_project):
    project_id, conn = memory_project

    with conn:
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('n0', 'doc_1', NULL, 'paragraph', 'Node 0', 0, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('n1', 'doc_1', NULL, 'paragraph', 'Original long content', 1, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('n2', 'doc_1', NULL, 'paragraph', 'Node 2', 2, 'current');"
        )

    upper, lower = split_node(conn, "n1", "Original long", "content")

    assert upper.id == "n1"
    assert upper.order_index == 1
    assert upper.embedding_status == "stale"

    assert lower.order_index == 2
    assert lower.embedding_status == "missing"

    cursor = conn.cursor()
    cursor.execute("SELECT id, order_index FROM nodes WHERE document_id = 'doc_1' ORDER BY order_index ASC;")
    all_nodes = cursor.fetchall()
    order_map = {r["id"]: r["order_index"] for r in all_nodes}

    assert order_map["n0"] == 0
    assert order_map["n1"] == 1
    assert order_map[lower.id] == 2
    assert order_map["n2"] == 3

def test_merge_combines_metadata_and_compacts_indices(memory_project):
    project_id, conn = memory_project

    with conn:
        conn.execute(
            "INSERT INTO schema_fields (id, project_id, field_slug, field_label, field_type, order_index) VALUES ('f_tags', ?, 'tags', 'Tags', 'array[string]', 0);",
            (project_id,),
        )
        conn.execute(
            "INSERT INTO schema_fields (id, project_id, field_slug, field_label, field_type, order_index) VALUES ('f_status', ?, 'status', 'Status', 'string', 1);",
            (project_id,),
        )

        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('lead', 'doc_1', NULL, 'paragraph', 'First part.', 0, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('succ', 'doc_1', NULL, 'paragraph', 'Second part.', 1, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('tail', 'doc_1', NULL, 'paragraph', 'Third part.', 2, 'current');"
        )

        conn.execute(
            "INSERT INTO node_metadata (node_id, field_id, field_value, user_edited) VALUES ('lead', 'f_tags', '[\"ai\", \"rag\"]', 0);"
        )
        conn.execute(
            "INSERT INTO node_metadata (node_id, field_id, field_value, user_edited) VALUES ('succ', 'f_tags', '[\"rag\", \"search\"]', 1);"
        )
        conn.execute(
            "INSERT INTO node_metadata (node_id, field_id, field_value, user_edited) VALUES ('lead', 'f_status', '\"draft\"', 0);"
        )
        conn.execute(
            "INSERT INTO node_metadata (node_id, field_id, field_value, user_edited) VALUES ('succ', 'f_status', '\"review\"', 0);"
        )

    merged = merge_nodes(conn, "lead")
    assert merged.text_content == "First part.\n\nSecond part."
    assert merged.order_index == 0

    cursor = conn.cursor()
    cursor.execute("SELECT id, order_index FROM nodes WHERE document_id = 'doc_1' ORDER BY order_index ASC;")
    remaining = cursor.fetchall()
    assert len(remaining) == 2
    assert remaining[0]["id"] == "lead"
    assert remaining[0]["order_index"] == 0
    assert remaining[1]["id"] == "tail"
    assert remaining[1]["order_index"] == 1

    cursor.execute("SELECT field_id, field_value, user_edited FROM node_metadata WHERE node_id = 'lead';")
    meta = {r["field_id"]: (r["field_value"], r["user_edited"]) for r in cursor.fetchall()}

    # List deduplication
    tags_val = json.loads(meta["f_tags"][0])
    assert set(tags_val) == {"ai", "rag", "search"}
    assert meta["f_tags"][1] == 1

    # Conflict on scalar string
    status_parsed = json.loads(meta["f_status"][0])
    assert status_parsed["has_conflict"] is True

def test_promote_node_reparents_siblings_and_cascades_stale(memory_project):
    project_id, conn = memory_project

    with conn:
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('p_target', 'doc_1', NULL, 'paragraph', 'Promote Me', 0, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('c1', 'doc_1', NULL, 'paragraph', 'Child 1', 1, 'current');"
        )
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('c2', 'doc_1', NULL, 'paragraph', 'Child 2', 2, 'current');"
        )

    promoted = promote_node_to_header(conn, "p_target")
    assert promoted.node_type == "header"
    assert promoted.embedding_status == "stale"

    cursor = conn.cursor()
    cursor.execute("SELECT id, parent_id, embedding_status FROM nodes WHERE id IN ('c1', 'c2');")
    children = {r["id"]: (r["parent_id"], r["embedding_status"]) for r in cursor.fetchall()}

    assert children["c1"][0] == "p_target"
    assert children["c1"][1] == "stale"
    assert children["c2"][0] == "p_target"
    assert children["c2"][1] == "stale"

def test_detach_selection_to_header(memory_project):
    project_id, conn = memory_project

    with conn:
        conn.execute(
            "INSERT INTO nodes (id, document_id, parent_id, node_type, text_content, order_index, embedding_status) VALUES ('raw_node', 'doc_1', NULL, 'paragraph', 'Before text. [Extracted Header] After text.', 0, 'current');"
        )

    text = "Before text. [Extracted Header] After text."
    start = text.index("[Extracted Header]")
    end = start + len("[Extracted Header]")

    result_nodes = detach_selection_to_header(conn, "raw_node", start, end)
    assert len(result_nodes) == 3

    assert result_nodes[0].text_content == "Before text."
    assert result_nodes[0].node_type == "paragraph"
    assert result_nodes[0].order_index == 0

    assert result_nodes[1].text_content == "[Extracted Header]"
    assert result_nodes[1].node_type == "header"
    assert result_nodes[1].order_index == 1

    assert result_nodes[2].text_content == "After text."
    assert result_nodes[2].node_type == "paragraph"
    assert result_nodes[2].parent_id == result_nodes[1].id
    assert result_nodes[2].order_index == 2

def test_document_and_node_api_routes(memory_project):
    project_id, conn = memory_project

    # Test upload via multipart endpoint
    file_bytes = b"# API Test Document\n\nFirst chunk paragraph."
    response = client.post(
        f"/api/projects/{project_id}/documents/upload",
        files=[("files", ("test.md", file_bytes, "text/markdown"))],
    )
    assert response.status_code == 200
    docs = response.json()
    assert len(docs) == 1
    new_doc_id = docs[0]["id"]
    assert docs[0]["total_nodes"] == 2

    # Query documents
    get_resp = client.get(f"/api/projects/{project_id}/documents")
    assert get_resp.status_code == 200
    doc_list = get_resp.json()
    assert len(doc_list) >= 2

    # Query nodes
    nodes_resp = client.get(f"/api/projects/{project_id}/nodes?document_id={new_doc_id}")
    assert nodes_resp.status_code == 200
    nodes = nodes_resp.json()
    assert len(nodes) == 2
    paragraph_node = [n for n in nodes if n["node_type"] == "paragraph"][0]

    # Test split endpoint
    split_resp = client.post(
        f"/api/projects/{project_id}/nodes/{paragraph_node['id']}/split",
        json={"top_text": "First chunk", "bottom_text": "paragraph."},
    )
    assert split_resp.status_code == 200
    split_data = split_resp.json()
    assert len(split_data) == 2
    assert split_data[0]["text_content"] == "First chunk"
    assert split_data[1]["text_content"] == "paragraph."